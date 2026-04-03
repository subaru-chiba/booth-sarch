from contextlib import asynccontextmanager
from typing import Optional
import io
import json
from datetime import datetime

import aiosqlite
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from fastapi import FastAPI, Form, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from .database import get_db, init_db
from .scraper import (
    CATEGORIES, SORT_OPTIONS,
    fetch_product, parse_item_id, search_category,
)

scheduler = AsyncIOScheduler()
templates = Jinja2Templates(directory="templates")


@asynccontextmanager
async def lifespan(app: FastAPI):
    await init_db()
    scheduler.add_job(refresh_all_products, "interval", hours=6, id="auto_refresh")
    scheduler.start()
    yield
    scheduler.shutdown()


app = FastAPI(title="Booth Sales Trend Tracker", lifespan=lifespan)
app.mount("/static", StaticFiles(directory="static"), name="static")


# ---------------------------------------------------------------------------
# Background task
# ---------------------------------------------------------------------------

async def refresh_all_products():
    db = await get_db()
    try:
        cursor = await db.execute("SELECT id, booth_item_id FROM products")
        products = await cursor.fetchall()
        for row in products:
            try:
                info = await fetch_product(row["booth_item_id"])
                await db.execute(
                    """INSERT INTO snapshots (product_id, sold_count, wish_count, price)
                       VALUES (?, ?, ?, ?)""",
                    (row["id"], info.sold_count, info.wish_count, info.price),
                )
                await db.execute(
                    """UPDATE products SET name=?, shop_name=?, shop_url=?, image_url=?, category=?
                       WHERE id=?""",
                    (info.name, info.shop_name, info.shop_url, info.image_url, info.category, row["id"]),
                )
            except Exception:
                pass
        await db.commit()
    finally:
        await db.close()


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    """Category selection home page."""
    return templates.TemplateResponse(request, "index.html", {
        "categories": CATEGORIES,
    })


@app.get("/browse", response_class=HTMLResponse)
async def browse(
    request: Request,
    category_id: Optional[int] = Query(default=None),
    sort: str = Query(default="new"),
    page: int = Query(default=1),
):
    """Browse Booth products by category, fetched live from Booth API."""
    try:
        products, total_pages = await search_category(
            category_id=category_id, sort=sort, page=page
        )
    except Exception as e:
        products, total_pages = [], 1
        error = str(e)
    else:
        error = None

    # Lookup category name
    cat_name = "すべて"
    for c in CATEGORIES:
        if c["id"] == category_id:
            cat_name = c["name"]
            break

    # Check which products are already tracked
    db = await get_db()
    try:
        cursor = await db.execute("SELECT booth_item_id FROM products")
        tracked_ids = {row[0] for row in await cursor.fetchall()}
    finally:
        await db.close()

    return templates.TemplateResponse(request, "browse.html", {
        "products": products,
        "categories": CATEGORIES,
        "sort_options": SORT_OPTIONS,
        "current_category_id": category_id,
        "current_category_name": cat_name,
        "current_sort": sort,
        "current_page": page,
        "total_pages": total_pages,
        "tracked_ids": tracked_ids,
        "error": error,
    })


@app.post("/track")
async def track_product(
    booth_item_id: int = Form(...),
    redirect_to: str = Form(default="/browse"),
):
    """Start tracking a product (add to DB with initial snapshot)."""
    try:
        info = await fetch_product(booth_item_id)
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))

    db = await get_db()
    try:
        await db.execute(
            """INSERT OR IGNORE INTO products
               (booth_item_id, name, shop_name, shop_url, image_url, category)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (info.booth_item_id, info.name, info.shop_name,
             info.shop_url, info.image_url, info.category),
        )
        cursor = await db.execute(
            "SELECT id FROM products WHERE booth_item_id = ?", (booth_item_id,)
        )
        row = await cursor.fetchone()
        await db.execute(
            """INSERT INTO snapshots (product_id, sold_count, wish_count, price)
               VALUES (?, ?, ?, ?)""",
            (row["id"], info.sold_count, info.wish_count, info.price),
        )
        await db.commit()
    finally:
        await db.close()

    return RedirectResponse(url=redirect_to, status_code=303)


@app.get("/tracked", response_class=HTMLResponse)
async def tracked(
    request: Request,
    category: Optional[str] = Query(default=None),
    sort: str = Query(default="sold_30d"),
):
    """List of products being tracked with trend data."""
    db = await get_db()
    try:
        cursor = await db.execute(
            "SELECT DISTINCT category FROM products WHERE category IS NOT NULL ORDER BY category"
        )
        categories = [row[0] for row in await cursor.fetchall()]

        query = """
            SELECT
                p.id, p.booth_item_id, p.name, p.shop_name, p.image_url, p.category,
                s_latest.sold_count, s_latest.wish_count, s_latest.price, s_latest.fetched_at,
                (
                    s_latest.sold_count - COALESCE(
                        (SELECT sold_count FROM snapshots
                         WHERE product_id = p.id
                           AND fetched_at <= datetime('now', '-30 days')
                         ORDER BY fetched_at DESC LIMIT 1),
                        (SELECT sold_count FROM snapshots
                         WHERE product_id = p.id
                         ORDER BY fetched_at ASC LIMIT 1)
                    )
                ) AS sold_30d
            FROM products p
            LEFT JOIN snapshots s_latest ON s_latest.id = (
                SELECT id FROM snapshots WHERE product_id = p.id ORDER BY fetched_at DESC LIMIT 1
            )
        """
        params: list = []
        if category:
            query += " WHERE p.category = ?"
            params.append(category)

        order = {
            "sold_30d":    "sold_30d DESC NULLS LAST",
            "sold_total":  "s_latest.sold_count DESC NULLS LAST",
            "wish":        "s_latest.wish_count DESC NULLS LAST",
            "added":       "p.added_at DESC",
        }.get(sort, "sold_30d DESC NULLS LAST")
        query += f" ORDER BY {order}"

        cursor = await db.execute(query, params)
        products = [dict(row) for row in await cursor.fetchall()]
    finally:
        await db.close()

    return templates.TemplateResponse(request, "tracked.html", {
        "products": products,
        "categories": categories,
        "current_category": category,
        "current_sort": sort,
    })


@app.get("/products/{item_id}", response_class=HTMLResponse)
async def product_detail(request: Request, item_id: int):
    db = await get_db()
    try:
        cursor = await db.execute(
            "SELECT * FROM products WHERE booth_item_id = ?", (item_id,)
        )
        product = await cursor.fetchone()
        if not product:
            raise HTTPException(status_code=404, detail="商品が見つかりません")
        product = dict(product)

        cursor = await db.execute(
            """SELECT sold_count, wish_count, price, fetched_at
               FROM snapshots WHERE product_id = ?
               ORDER BY fetched_at ASC""",
            (product["id"],),
        )
        snapshots = [dict(row) for row in await cursor.fetchall()]
    finally:
        await db.close()

    chart_labels = [s["fetched_at"] for s in snapshots]
    chart_sold   = [s["sold_count"] for s in snapshots]
    chart_wish   = [s["wish_count"] for s in snapshots]
    chart_price  = [s["price"] for s in snapshots]

    return templates.TemplateResponse(request, "product.html", {
        "product": product,
        "snapshots": snapshots,
        "chart_labels": json.dumps(chart_labels),
        "chart_sold":   json.dumps(chart_sold),
        "chart_wish":   json.dumps(chart_wish),
        "chart_price":  json.dumps(chart_price),
    })


@app.post("/products/{item_id}/refresh")
async def refresh_product(item_id: int):
    db = await get_db()
    try:
        cursor = await db.execute(
            "SELECT id FROM products WHERE booth_item_id = ?", (item_id,)
        )
        row = await cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404)
        product_id = row["id"]
        info = await fetch_product(item_id)
        await db.execute(
            """INSERT INTO snapshots (product_id, sold_count, wish_count, price)
               VALUES (?, ?, ?, ?)""",
            (product_id, info.sold_count, info.wish_count, info.price),
        )
        await db.execute(
            """UPDATE products SET name=?, shop_name=?, shop_url=?, image_url=?, category=?
               WHERE id=?""",
            (info.name, info.shop_name, info.shop_url, info.image_url, info.category, product_id),
        )
        await db.commit()
    finally:
        await db.close()
    return RedirectResponse(url=f"/products/{item_id}", status_code=303)


@app.post("/products/{item_id}/delete")
async def delete_product(item_id: int):
    db = await get_db()
    try:
        await db.execute("DELETE FROM products WHERE booth_item_id = ?", (item_id,))
        await db.commit()
    finally:
        await db.close()
    return RedirectResponse(url="/tracked", status_code=303)


@app.get("/export/xlsx")
async def export_xlsx(category: Optional[str] = Query(default=None)):
    db = await get_db()
    try:
        query = """
            SELECT
                p.booth_item_id, p.name, p.category, p.shop_name,
                s_latest.price, s_latest.sold_count, s_latest.wish_count, s_latest.fetched_at,
                (
                    s_latest.sold_count - COALESCE(
                        (SELECT sold_count FROM snapshots
                         WHERE product_id = p.id
                           AND fetched_at <= datetime('now', '-30 days')
                         ORDER BY fetched_at DESC LIMIT 1),
                        (SELECT sold_count FROM snapshots
                         WHERE product_id = p.id
                         ORDER BY fetched_at ASC LIMIT 1)
                    )
                ) AS sold_30d
            FROM products p
            LEFT JOIN snapshots s_latest ON s_latest.id = (
                SELECT id FROM snapshots WHERE product_id = p.id ORDER BY fetched_at DESC LIMIT 1
            )
        """
        params: list = []
        if category:
            query += " WHERE p.category = ?"
            params.append(category)
        query += " ORDER BY sold_30d DESC NULLS LAST"
        cursor = await db.execute(query, params)
        rows = [dict(r) for r in await cursor.fetchall()]
    finally:
        await db.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "売れ行きトレンド"

    header_fill = PatternFill(fill_type="solid", fgColor="1a2a3a")
    header_font = Font(bold=True, color="5d9ee8")

    headers = [
        ("商品ID", 12), ("商品名", 40), ("カテゴリ", 18), ("ショップ名", 22),
        ("価格", 10), ("直近30日売上", 14), ("累計売上", 12),
        ("ウィッシュ数", 12), ("最終取得日時", 20),
    ]
    for col, (label, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = width

    for row in rows:
        ws.append([
            row["booth_item_id"], row["name"], row["category"] or "",
            row["shop_name"] or "", row["price"], row["sold_30d"],
            row["sold_count"], row["wish_count"],
            row["fetched_at"][:16] if row["fetched_at"] else "",
        ])

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"booth_trend_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
